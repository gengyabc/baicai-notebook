import csv
import json
import subprocess
from pathlib import Path

from template_gen.generate_fill_data import (
    generate_fill_data,
    load_placeholder_descriptions,
    normalize_placeholder_key,
)
from template_gen.export_placeholder_csv import export_placeholder_csv
from template_gen.import_placeholder_csv import import_placeholder_csv

SECRET_BINDING_FIELD = "secret_name"

REPO_ROOT = Path(__file__).resolve().parent.parent
OPENCODE_DIR = REPO_ROOT / ".opencode"
HARNESS = OPENCODE_DIR / "scripts" / "run-secure-action-harness.mjs"


def _write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def _bun(*args):
    result = subprocess.run(
        ["bun", "run", str(HARNESS), *args],
        capture_output=True,
        text=True,
        cwd=str(OPENCODE_DIR),
        timeout=30,
    )
    return result


class TestBatch1FillDocxWorkflowSeparation:
    def test_fill_docx_separates_llm_and_local_sensitive_path(self):
        text = (REPO_ROOT / ".opencode" / "workflows" / "fill-docx.md").read_text()
        assert "normal LLM fill path" in text
        assert "local-only sensitive fill path" in text
        assert "must never be sent to the model" in text

    def test_fill_docx_states_sensitive_values_must_never_be_sent_to_model(self):
        text = (REPO_ROOT / ".opencode" / "workflows" / "fill-docx.md").read_text()
        assert "sensitive values must never be sent to the model" in text

    def test_fill_docx_mentions_secret_name_binding(self):
        text = (REPO_ROOT / ".opencode" / "workflows" / "fill-docx.md").read_text()
        assert "secret_name" in text

    def test_fill_docx_mentions_opaque_fill_step(self):
        text = (REPO_ROOT / ".opencode" / "workflows" / "fill-docx.md").read_text()
        assert "Opaque Sensitive Fill" in text

    def test_fill_docx_mentions_secure_action(self):
        text = (REPO_ROOT / ".opencode" / "workflows" / "fill-docx.md").read_text()
        assert "fill-sensitive-fill-data.mjs" in text or "Opaque Sensitive Fill" in text

    def test_fill_docx_llm_must_leave_secret_fields_empty(self):
        text = (REPO_ROOT / ".opencode" / "workflows" / "fill-docx.md").read_text()
        assert "secret_name" in text
        assert "leave" in text.lower() or "empty" in text.lower()

class TestBatch1PermissionBoundary:
    def test_opencode_json_denies_raw_secret_access(self):
        config = json.loads((REPO_ROOT / "opencode.json").read_text())
        perm = config["permission"]
        bash_denies = list(perm.get("bash", {}).keys())
        assert any("cross-keychain" in k for k in bash_denies)
        assert any("keytar" in k for k in bash_denies)


class TestCsvRoundtripPreservesSecretMetadata:
    def test_export_includes_secret_name_column(self, tmp_path: Path):
        placeholders_json = tmp_path / "placeholders.json"
        out_csv = tmp_path / "out" / "descriptions.csv"
        _write_json(
            placeholders_json,
            {
                "placeholders": [
                    {"placeholder": "{{ phone }}", "description": "手机号码", SECRET_BINDING_FIELD: "PHONE"},
                    {"placeholder": "{{ project_name }}", "description": "项目名称"},
                ]
            },
        )

        export_placeholder_csv(str(placeholders_json), str(out_csv))

        rows = _read_csv_rows(out_csv)
        assert rows[0]["secret_name"] == "PHONE"
        assert rows[1]["secret_name"] == ""

    def test_import_preserves_secret_name(self, tmp_path: Path):
        source_csv = tmp_path / "descriptions.csv"
        out_json = tmp_path / "out" / "descriptions.json"
        source_csv.write_text(
            "placeholder,description,secret_name\n{{ phone }},手机号码,PHONE\n{{ project_name }},项目名称,\n",
            encoding="utf-8",
        )

        import_placeholder_csv(str(source_csv), str(out_json))

        payload = _read_json(out_json)
        assert payload["placeholders"][0][SECRET_BINDING_FIELD] == "PHONE"
        assert SECRET_BINDING_FIELD not in payload["placeholders"][1] or payload["placeholders"][1][SECRET_BINDING_FIELD] == ""

    def test_roundtrip_preserves_secret_metadata(self, tmp_path: Path):
        placeholders_json = tmp_path / "placeholders.json"
        csv_path = tmp_path / "descriptions.csv"
        reimported_json = tmp_path / "descriptions.json"

        _write_json(
            placeholders_json,
            {
                "placeholders": [
                    {"placeholder": "{{ phone }}", "description": "手机号码", SECRET_BINDING_FIELD: "PHONE"},
                    {"placeholder": "{{ project_name }}", "description": "项目名称"},
                ]
            },
        )

        export_placeholder_csv(str(placeholders_json), str(csv_path))
        import_placeholder_csv(str(csv_path), str(reimported_json))

        payload = _read_json(reimported_json)
        assert payload["placeholders"][0][SECRET_BINDING_FIELD] == "PHONE"
        assert SECRET_BINDING_FIELD not in payload["placeholders"][1] or payload["placeholders"][1][SECRET_BINDING_FIELD] == ""

    def test_edit_mode_rebuild_preserves_secret_name_for_surviving_placeholders(self, tmp_path: Path):
        from docx import Document
        from template_gen.export_placeholder_csv import export_placeholder_csv

        placeholders_json = tmp_path / "temp" / "placeholders.json"
        out_csv = tmp_path / "output" / "descriptions.csv"
        template_docx = tmp_path / "output" / "template.docx"

        _write_json(
            placeholders_json,
            {
                "placeholders": [
                    {"location": "paragraphs[0]", "placeholder": "{{ phone }}", SECRET_BINDING_FIELD: "PHONE"},
                    {"location": "tables[0].rows[0].cells[0]", "placeholder": "{{ project_name }}"},
                ]
            },
        )

        doc = Document()
        doc.add_paragraph("{{ phone }}")
        doc.add_paragraph("{{ project_name }}")
        template_docx.parent.mkdir(parents=True, exist_ok=True)
        doc.save(template_docx)

        export_placeholder_csv(
            str(placeholders_json),
            str(out_csv),
            edit=True,
            template_docx_path=str(template_docx),
            placeholders_output_path=str(placeholders_json),
        )

        refreshed = _read_json(placeholders_json)
        phone_entry = next(p for p in refreshed["placeholders"] if p["placeholder"] == "{{ phone }}")
        assert phone_entry[SECRET_BINDING_FIELD] == "PHONE"

    def test_edit_mode_rebuild_drops_secret_name_for_removed_placeholders(self, tmp_path: Path):
        from docx import Document
        from template_gen.export_placeholder_csv import export_placeholder_csv

        placeholders_json = tmp_path / "temp" / "placeholders.json"
        out_csv = tmp_path / "output" / "descriptions.csv"
        template_docx = tmp_path / "output" / "template.docx"

        _write_json(
            placeholders_json,
            {
                "placeholders": [
                    {"location": "paragraphs[0]", "placeholder": "{{ phone }}", SECRET_BINDING_FIELD: "PHONE"},
                    {"location": "paragraphs[1]", "placeholder": "{{ old_field }}"},
                ]
            },
        )

        doc = Document()
        doc.add_paragraph("{{ phone }}")
        template_docx.parent.mkdir(parents=True, exist_ok=True)
        doc.save(template_docx)

        export_placeholder_csv(
            str(placeholders_json),
            str(out_csv),
            edit=True,
            template_docx_path=str(template_docx),
            placeholders_output_path=str(placeholders_json),
        )

        refreshed = _read_json(placeholders_json)
        assert not any(p["placeholder"] == "{{ old_field }}" for p in refreshed["placeholders"])


class TestGenerateFillDataWithSecretBinding:
    def test_generate_fill_data_ignores_secret_name_for_fill_values(self, tmp_path: Path):
        repo = tmp_path / "repo"
        (repo / ".opencode").mkdir(parents=True, exist_ok=True)
        (repo / ".opencode" / "vault-config.json").write_text(
            '{"vaultRoot": "workbook"}', encoding="utf-8"
        )
        (repo / "workbook").mkdir(parents=True, exist_ok=True)

        source = repo / "descriptions.json"
        output = repo / "out" / "fill_data.json"
        _write_json(
            source,
            {
                "placeholders": [
                    {"placeholder": "{{ phone }}", "description": "手机号码", SECRET_BINDING_FIELD: "PHONE"},
                    {"placeholder": "{{ project_name }}", "description": "项目名称"},
                ]
            },
        )

        generate_fill_data(str(source), str(output), repo_root=repo)

        loaded = _read_json(output)
        assert loaded == {"phone": "", "project_name": ""}


class TestJsSensitiveFillDataInjection:
    def test_js_script_fills_secret_backed_fields_in_fill_data(self, tmp_path: Path):
        repo = tmp_path / "repo"
        (repo / ".opencode").mkdir(parents=True, exist_ok=True)
        (repo / ".opencode" / "vault-config.json").write_text(
            '{"vaultRoot": "workbook"}', encoding="utf-8"
        )
        (repo / "workbook").mkdir(parents=True, exist_ok=True)

        descriptions = repo / ".temp" / "职业教育优质教材推荐书" / "temp-v1" / "descriptions.json"
        fill_data = repo / ".temp" / "职业教育优质教材推荐书" / "temp-v1" / "fill_data.json"
        descriptions.parent.mkdir(parents=True, exist_ok=True)
        _write_json(
            descriptions,
            {
                "placeholders": [
                    {"placeholder": "{{ phone }}", "description": "联系人手机", "secret_name": "PHONE"},
                    {"placeholder": "{{ email }}", "description": "电子邮箱"},
                ]
            },
        )
        _write_json(fill_data, {"phone": "", "email": "test@example.invalid"})

        _bun("seed-secret", "PHONE", "test-phone-secret-value")
        try:
            result = subprocess.run(
                [
                    "bun",
                    "run",
                    str(REPO_ROOT / ".opencode" / "scripts" / "fill-sensitive-fill-data.mjs"),
                    "--descriptions",
                    str(descriptions),
                    "--fill-data",
                    str(fill_data),
                ],
                capture_output=True,
                text=True,
                cwd=str(REPO_ROOT),
                timeout=30,
            )
            assert result.returncode == 0, result.stderr
            assert "test-phone-secret-value" not in result.stdout

            loaded = _read_json(fill_data)
            assert loaded["phone"] == "test-phone-secret-value"
            assert loaded["email"] == "test@example.invalid"
        finally:
            _bun("remove-secret", "PHONE")


class TestDocFillInfrastructureReadiness:
    def test_allowlist_schema_supports_new_services(self):
        allowlist = json.loads(
            (OPENCODE_DIR / "plugin-allowlist.json").read_text()
        )
        assert "services" in allowlist
        assert isinstance(allowlist["services"], dict)

    def test_plugin_handler_dispatch_is_static(self):
        source = (OPENCODE_DIR / "plugins" / "secure-plugin-access.ts").read_text()
        assert "HANDLERS" in source
        assert "require(" not in source or "createRequire" in source

    def test_plugin_validation_order_preserves_secret_safety(self):
        source = (OPENCODE_DIR / "plugins" / "secure-plugin-access.ts").read_text()
        assert "unknown-service" in source
        assert "unknown-operation" in source
        assert "invalid-args" in source
        assert "secret-missing" in source
